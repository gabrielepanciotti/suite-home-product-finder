"""DEMO: confronto con immagine prodotto inline (solo top picks).

Scarica le foto prodotto, le rimpicciolisce e le incorpora in una colonna "Foto"
dell'Excel. Versione dimostrativa su pochi box per validare il look prima di
estenderlo a tutti.
"""
import io
from pathlib import Path
import httpx
from PIL import Image as PILImage
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
     "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
     "Accept-Language": "it-IT,it;q=0.9"}

# coll, box, dim, cap, eur, ruote, disp, link, img_url
ROWS = [
    ("Sotto il letto", "IKEA SAMLA 55L", "78x56x18", 55, 13.00, "No", "Disponibile",
     "https://www.ikea.com/it/it/p/samla-contenitore-trasparente-90102971/",
     "https://www.ikea.com/it/it/images/products/samla-contenitore-trasparente__0711155_pe728023_s5.jpg"),
    ("Armadio/mensola", "Maury's porta indumenti 60x41x40", "60x41x40", 74, 9.90, "Ruote", "Disponibile",
     "https://maurysonline.it/box-contenitore-con-ruote-dimensioni-60x41x40cm-e-coperchio-bianco",
     "https://cdn.maurysonline.it/immagini/6075081_0438026-box-porta-indumenti-con-ruote-dimensioni-60x41x40cm-e-coperchio-bianco-main_510.jpeg"),
    ("Armadio/mensola", "JYSK Basic Box 80L", "60x40x43", 80, 12.50, "No", "Disponibile",
     "https://jysk.it/organizzazione/contenitori-e-cesti/contenitori-plastica/contenitore-basic-box-80-l-con-coperchio-trasparente",
     "https://cdn1.jysk.com/getimage/wd3.large/238798"),
    ("Cantina/garage", "IKEA SAMLA 130L", "79x57x43", 130, 20.00, "No", "Disponibile",
     "https://www.ikea.com/it/it/p/samla-contenitore-con-coperchio-trasparente-s29440819/",
     "https://www.ikea.com/it/it/images/products/samla-contenitore-con-coperchio-trasparente__0711819_pe728498_s5.jpg"),
    ("Cantina/garage", "Iris Ohyama 80L (set 4)", "84x42.5x33", 80, 23.51, "No", "Disponibile",
     "https://www.amazon.it/dp/B0CDCP8FB7",
     "https://m.media-amazon.com/images/I/71ggSp38aTL._AC_SL1500_.jpg"),
    ("Cantina/garage", "Tontarelli Box One 140L", "79x52x52", 140, 29.00, "No", "Disponibile",
     "https://www.orizzonteshop.it/products/tontarelli-boxone-storage-ccoperchio-8035940999-tontarelli-ean-8009404212080",
     "http://www.orizzonteshop.it/cdn/shop/products/immagine-1-tontarelli-boxone-storage-ccoperchio-8035940999-tontarelli-ean-8009404212080.jpg?height=1200&width=1200"),
]

THUMB_H = 80  # px

def fetch_thumb(url):
    with httpx.Client(headers=H, timeout=30, follow_redirects=True) as c:
        data = c.get(url).content
    im = PILImage.open(io.BytesIO(data)).convert("RGB")
    w, h = im.size
    nw = int(w * THUMB_H / h)
    im = im.resize((nw, THUMB_H), PILImage.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return buf, nw

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Confronto (demo foto)"
HDR = ["Collocazione", "Foto", "Box", "Dim (cm)", "Cap (L)", "Prezzo (EUR)",
       "EUR/Litro", "Ruote", "Disponibilita", "Link"]
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
title = ws.cell(1, 1, "DEMO - confronto scatole vestiti con foto prodotto inline (top picks)")
title.font = Font(bold=True, size=11)
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
for j, h in enumerate(HDR, 1):
    c = ws.cell(3, j, h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border
collfill = {"Sotto il letto": "E8F0FE", "Armadio/mensola": "E6F4EA", "Cantina/garage": "FEF7E0"}
r = 4
for coll, box, dim, cap, eur, ruote, disp, link, img_url in ROWS:
    epl = round(eur / cap, 2)
    ws.cell(r, 1, coll).fill = PatternFill("solid", fgColor=collfill[coll])
    ws.cell(r, 3, box).font = Font(bold=True)
    ws.cell(r, 4, dim)
    ws.cell(r, 5, cap)
    ws.cell(r, 6, eur).number_format = "0.00"
    ws.cell(r, 7, epl).number_format = "0.00"
    ws.cell(r, 8, ruote)
    ws.cell(r, 9, disp)
    lk = ws.cell(r, 10, "apri")
    lk.hyperlink = link
    lk.font = Font(color="0563C1", underline="single")
    for j in range(1, 11):
        ws.cell(r, j).border = border
        ws.cell(r, j).alignment = Alignment(vertical="center", wrap_text=(j == 1))
    try:
        buf, nw = fetch_thumb(img_url)
        xi = XLImage(buf)
        xi.height = THUMB_H
        xi.width = nw
        ws.add_image(xi, f"B{r}")
        ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, nw / 7 + 2)
        print("img OK", box, nw, "x", THUMB_H)
    except Exception as e:
        ws.cell(r, 2, "(no img)")
        print("img ERR", box, type(e).__name__, str(e)[:50])
    ws.row_dimensions[r].height = THUMB_H * 0.78  # px -> points approx
    r += 1
ws.freeze_panes = "A4"
for j, w in enumerate([16, 16, 30, 14, 7, 11, 9, 8, 16, 7], 1):
    if get_column_letter(j) != "B":
        ws.column_dimensions[get_column_letter(j)].width = w
out = Path(__file__).resolve().parents[1] / "vault" / "prodotti" / "contenitori" / "scatole-vestiti-confronto-DEMO-foto.xlsx"
wb.save(out)
print("OK ->", out)
