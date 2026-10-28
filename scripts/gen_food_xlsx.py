"""RICERCA A - contenitori ALIMENTARI (dispensa/frigo). Excel con foto inline.

Baseline = prodotti fotografati a Maxi (Ma Maison, Maison du Monde, Frigochef) +
alternative LIVE Action (Mepal/Curver/Tefal/barattoli) e IKEA (365+/PRUTA/KORKEN).
Distinta dalla RICERCA B (scatole stoccaggio grandi -> scatole-vestiti-confronto).
Metrica = EUR/Litro. Foto Maxi = scatti locali in vault/raw; Action/IKEA = catalogo.
"""
import io
import os
from pathlib import Path
import httpx
from PIL import Image as PILImage
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36", "Accept-Language": "it-IT,it;q=0.9"}
RAW = str(Path(__file__).resolve().parents[1] / "vault" / "raw") + os.sep
AC = "https://asset.action.com/image/upload/t_digital_square/w_512/"
IK = "https://www.ikea.com/it/it/images/products/"

# gruppo, prodotto, negozio, cap(L), materiale, ermetico, eur, note, img, link
R = [
    # --- Dispensa / ermetici (secco, barattoli) ---
    ("Dispensa/ermetici", "Action barattolo ermetico 1,4L", "Action", 1.4, "Plastica", "Si", 2.72, "alt. economica Ma Maison",
     AC + "3218514_8712442227854-110_01_goq5xz.jpg", "https://www.action.com/it-it/p/3218514/barattolo-con-chiusura-ermetica/"),
    ("Dispensa/ermetici", "Action barattolo 0,8L", "Action", 0.8, "Plastica", "Si", 1.99, "",
     AC + "3218513_8712442227823-110_02_ev4jgi.jpg", "https://www.action.com/it-it/p/3218513/barattolo/"),
    ("Dispensa/ermetici", "Action contenitore vetro 1L", "Action", 1.0, "Vetro", "Si", 2.59, "vetro",
     AC + "3220679_3560231067373-110_05_iy2adl.jpg", "https://www.action.com/it-it/p/3220679/contenitore-in-vetro/"),
    ("Dispensa/ermetici", "Action barattolo vetro conserve 1L", "Action", 1.0, "Vetro", "Si", 2.99, "vetro, conserve",
     AC + "13792752.jpg", "https://www.action.com/it-it/p/2580021/barattolo-per-conserve/"),
    ("Dispensa/ermetici", "IKEA 365+ cibi secchi 1,3L", "IKEA", 1.3, "Plastica", "Si", 3.95, "il piu' venduto IKEA",
     IK + "ikea-365-contenitore-coperchio-cibi-secchi-trasparente-bianco__0711434_pe728250_s5.jpg",
     "https://www.ikea.com/it/it/p/ikea-365-contenitore-coperchio-cibi-secchi-trasparente-bianco-80066723/"),
    ("Dispensa/ermetici", "IKEA KORKEN barattolo vetro 1L", "IKEA", 1.0, "Vetro", "Si", 2.95, "vetro con guarnizione",
     IK + "korken-barattolo-con-coperchio-vetro-trasparente__1152252_pe885219_s5.jpg",
     "https://www.ikea.com/it/it/p/korken-barattolo-con-coperchio-vetro-trasparente-10541366/"),
    ("Dispensa/ermetici", "Ma Maison Airtight tondo 1L", "Maxi", 1.0, "Plastica", "Si", 4.99, "fotografato Maxi",
     RAW + "2026-06-14_158_foto.jpg", ""),
    ("Dispensa/ermetici", "Ma Maison Airtight 660ml", "Maxi", 0.66, "Plastica", "Si", 3.99, "fotografato Maxi",
     RAW + "2026-06-14_158_foto.jpg", ""),
    ("Dispensa/ermetici", "Ma Maison Clear&Fresh rett. 1L", "Maxi", 1.0, "Plastica", "Si", 3.99, "fotografato Maxi",
     RAW + "2026-06-14-ma-maison-clear-and-fresh-rettangolare-1l.jpg", ""),
    ("Dispensa/ermetici", "Ma Maison Clear&Fresh tondo 1L", "Maxi", 1.0, "Plastica", "Si", 3.99, "fotografato Maxi",
     RAW + "2026-06-14_159_foto.jpg", ""),
    ("Dispensa/ermetici", "Ma Maison Ying Yang 700ml", "Maxi", 0.7, "Plastica", "Si", 1.99, "fotografato Maxi",
     RAW + "2026-06-14_158_foto.jpg", ""),
    # --- Frigo / avanzi (micro/freezer) ---
    ("Frigo/avanzi", "IKEA 365+ contenitore 2,0L", "IKEA", 2.0, "Plastica", "Si", 2.50, "miglior EUR/L",
     IK + "ikea-365-contenitore-rettangolare-plastica__0711428_pe728244_s5.jpg",
     "https://www.ikea.com/it/it/p/ikea-365-contenitore-rettangolare-plastica-60359152/"),
    ("Frigo/avanzi", "Frigochef 1,25L", "Maxi", 1.25, "Plastica", "Coperchio", 2.59, "fotografato Maxi",
     RAW + "2026-06-14-scaffale-contenitori-maxi.jpg", ""),
    ("Frigo/avanzi", "Maison du Monde 1,6L clip", "Maxi", 1.6, "Plastica", "Si clip", 3.50, "fotografato Maxi",
     RAW + "2026-06-14-scaffale-contenitori-maxi.jpg", ""),
    ("Frigo/avanzi", "Action Mepal 1,5L", "Action", 1.5, "Plastica", "Si", 3.99, "brand premium",
     AC + "25077426.jpg", "https://www.action.com/it-it/p/3208098/contenitore-per-alimenti-mepal/"),
    ("Frigo/avanzi", "Action Curver 1,4L", "Action", 1.4, "Plastica", "Si", 3.79, "",
     AC + "17623549.jpg", "https://www.action.com/it-it/p/3002389/contenitore-per-alimenti-curver/"),
    ("Frigo/avanzi", "Action Mepal 0,75L", "Action", 0.75, "Plastica", "Si", 2.99, "brand premium",
     AC + "25077875.jpg", "https://www.action.com/it-it/p/3208100/contenitore-per-alimenti-mepal/"),
    ("Frigo/avanzi", "Maison du Monde 900ml", "Maxi", 0.9, "Plastica", "Coperchio", 4.90, "fotografato Maxi",
     RAW + "2026-06-14-scaffale-contenitori-maxi.jpg", ""),
    ("Frigo/avanzi", "Maison du Monde 1,2L", "Maxi", 1.2, "Plastica", "Coperchio", 5.90, "fotografato Maxi",
     RAW + "2026-06-14-scaffale-contenitori-maxi.jpg", ""),
    ("Frigo/avanzi", "Action Mepal 0,25L", "Action", 0.25, "Plastica", "Si", 1.88, "brand premium",
     AC + "25077236.jpg", "https://www.action.com/it-it/p/3208097/contenitore-per-alimenti-mepal/"),
    ("Frigo/avanzi", "Maison du Monde 400ml", "Maxi", 0.4, "Plastica", "Coperchio", 2.99, "fotografato Maxi",
     RAW + "2026-06-14-scaffale-contenitori-maxi.jpg", ""),
    # --- Set / multipack (EUR/L solo se capacita nota) ---
    ("Set/multipack", "Maison du Monde set 4x0,5L", "Maxi", 2.0, "Plastica", "Coperchio", 2.99, "4 pz = 2L tot",
     RAW + "2026-06-14-scaffale-contenitori-maxi.jpg", ""),
    ("Set/multipack", "Action Tefal 1L x2", "Action", 2.0, "Plastica", "Si", 6.95, "2 pz = 2L tot",
     AC + "28018846.jpg", "https://www.action.com/it-it/p/3208738/contenitori-per-alimenti-tefal/"),
    ("Set/multipack", "IKEA PRUTA set 17pz", "IKEA", None, "Plastica", "Si", 4.95, "17 pz misti, frigo/freezer",
     IK + "pruta-set-17-contenitori-alim-con-coperch-verde__1383829_pe962689_s5.jpg",
     "https://www.ikea.com/it/it/p/pruta-set-17-contenitori-alim-con-coperch-verde-10597099/"),
    ("Set/multipack", "Action contenitori impilabili 3pz", "Action", None, "Plastica", "Coperchio", 5.95, "3 pz",
     AC + "3218110_8721037720373-110_01_nujg9n.jpg", "https://www.action.com/it-it/p/3218110/contenitori-impilabili/"),
    ("Set/multipack", "IKEA 365+ set 6 contenitori", "IKEA", None, "Plastica", "Si", 17.00, "6 pz misti",
     IK + "ikea-365-set-6-contenit-alim-con-coperchio-plastica__1109095_pe869821_s5.jpg",
     "https://www.ikea.com/it/it/p/ikea-365-set-6-contenit-alim-con-coperchio-plastica-80526452/"),
]
best = {"Action barattolo ermetico 1,4L", "IKEA 365+ contenitore 2,0L", "Maison du Monde set 4x0,5L"}

THUMB_H = 72


def fetch_thumb(src, client):
    if os.path.exists(src):
        data = open(src, "rb").read()
    else:
        data = client.get(src).content
    im = PILImage.open(io.BytesIO(data)).convert("RGB")
    w, h = im.size
    nw = max(1, int(w * THUMB_H / h))
    im = im.resize((nw, THUMB_H), PILImage.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return buf, nw


wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Confronto alimentari"
HDR = ["Gruppo", "Foto", "Prodotto", "Negozio", "Cap (L)", "Materiale", "Ermetico",
       "Prezzo (EUR)", "EUR/Litro", "Note", "Link"]
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
title = ws.cell(1, 1, "RICERCA A - contenitori ALIMENTARI (dispensa/frigo). Baseline = "
                      "fotografati a Maxi (Ma Maison/Maison du Monde/Frigochef); alternative "
                      "Action e IKEA verificate live 14/06. EUR/Litro = valore (verde = top per "
                      "gruppo). NB: scatole stoccaggio grandi = ricerca separata (scatole-vestiti-confronto).")
title.font = Font(bold=True, size=11)
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
title.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[1].height = 44
for j, h in enumerate(HDR, 1):
    c = ws.cell(3, j, h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border
grpfill = {"Dispensa/ermetici": "E8F0FE", "Frigo/avanzi": "E6F4EA", "Set/multipack": "FEF7E0"}
imgcol_w = 12.0
r = 4
with httpx.Client(headers=UA, timeout=30, follow_redirects=True) as client:
    for grp, nome, neg, cap, mat, erm, eur, note, img, link in R:
        epl = round(eur / cap, 2) if cap else None
        ws.cell(r, 1, grp).fill = PatternFill("solid", fgColor=grpfill[grp])
        ws.cell(r, 3, nome).font = Font(bold=(nome in best))
        ws.cell(r, 4, neg)
        ws.cell(r, 5, cap if cap else "-")
        ws.cell(r, 6, mat)
        ws.cell(r, 7, erm)
        ws.cell(r, 8, eur).number_format = "0.00"
        epl_c = ws.cell(r, 9, epl if epl is not None else "-")
        if epl is not None:
            epl_c.number_format = "0.00"
        ws.cell(r, 10, note)
        lk = ws.cell(r, 11, "apri" if link else neg)
        if link:
            lk.hyperlink = link
            lk.font = Font(color="0563C1", underline="single")
        if nome in best and epl is not None:
            epl_c.fill = PatternFill("solid", fgColor="C6EFCE")
        for j in range(1, 12):
            ws.cell(r, j).border = border
            ws.cell(r, j).alignment = Alignment(vertical="center", wrap_text=(j in (1, 10)))
        try:
            buf, nw = fetch_thumb(img, client)
            xi = XLImage(buf)
            xi.height = THUMB_H
            xi.width = nw
            ws.add_image(xi, f"B{r}")
            imgcol_w = max(imgcol_w, nw / 7 + 2)
            print("img OK ", nome)
        except Exception as e:
            ws.cell(r, 2, "(no img)").alignment = Alignment(horizontal="center", vertical="center")
            print("no img ", nome, type(e).__name__, str(e)[:40])
        ws.row_dimensions[r].height = THUMB_H * 0.78
        r += 1
ws.freeze_panes = "A4"
ws.auto_filter.ref = f"A3:K{r - 1}"
W = [16, None, 30, 9, 7, 11, 9, 11, 9, 24, 8]
for j, w in enumerate(W, 1):
    col = get_column_letter(j)
    ws.column_dimensions[col].width = imgcol_w if col == "B" else w

ws2 = wb.create_sheet("Sintesi")
rows = [
    ("RACCOMANDAZIONE - contenitori alimentari", ""),
    ("", ""),
    ("Dispensa / ermetici", "Action barattolo ermetico 1,4L - 2,72 EUR (1,94 EUR/L) = molto piu' "
     "conveniente del Ma Maison Airtight 1L (4,99) e piu' grande. Per il vetro: Action contenitore "
     "vetro 1L 2,59 EUR o IKEA KORKEN 1L 2,95 EUR."),
    ("Frigo / avanzi", "IKEA 365+ 2,0L - 2,50 EUR (1,25 EUR/L) = miglior valore. Maison du Monde 1,6L "
     "clip (3,50) buono per il volume; Action Mepal/Curver se vuoi il brand."),
    ("Set / multipack", "Maison du Monde set 4x0,5L 2,99 EUR (1,50 EUR/L) ottimo. IKEA PRUTA 17pz 4,95 "
     "EUR per tanti piccoli (frigo/freezer)."),
    ("", ""),
    ("Verdetto Maxi", "I Ma Maison/Maison du Monde di Maxi NON sono i piu' convenienti: Action batte "
     "quasi sempre l'ermetico da dispensa, IKEA 365+ batte il frigo. Maxi ok per comprare subito senza spedizione."),
    ("Nota 2 ricerche", "Questa = contenitori alimentari piccoli (dispensa/frigo). Le scatole grandi da "
     "stoccaggio (vestiti/ingombranti) sono la ricerca separata 'scatole-vestiti-confronto'."),
]
ws2.cell(1, 1).font = Font(bold=True, size=12)
for i, (a, b) in enumerate(rows, 1):
    ws2.cell(i, 1, a).font = Font(bold=True)
    ws2.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 100

out = Path(__file__).resolve().parents[1] / "vault" / "prodotti" / "contenitori" / "contenitori-alimentari-confronto.xlsx"
wb.save(out)
print("OK ->", out, "|", len(R), "prodotti")
